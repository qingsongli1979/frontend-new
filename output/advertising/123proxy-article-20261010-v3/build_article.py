from pathlib import Path
from io import BytesIO
import json

from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
OUT = ROOT / '123Proxy_不限流量海外爬虫代理_广告图文_2026-10.docx'
FONT = 'Microsoft YaHei'
BLUE = '2457C5'
OFFERS = json.loads((ASSETS / 'live-offers.json').read_text(encoding='utf-8-sig'))['offers']


def price(kind, amount):
    matches = [offer for offer in OFFERS if offer['chargeType'] == kind and offer['amount'] == amount]
    assert len(matches) == 1, (kind, amount)
    offer = matches[0]
    assert offer['unit'] == 'month'
    return round(offer['price'] * offer['discount'])


def money(value):
    return f'¥{value:,}'


# Keep the generated price artwork tied to the verified public quote snapshot.
assert [price('tunnelIp', amount) for amount in (5, 50, 200)] == [270, 2700, 9000]
assert price('durationIp', 5) == 600


def run_font(run, size=None, bold=None, color='000000'):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn(f'w:{key}'), FONT)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.68)
section.bottom_margin = Inches(.66)
section.left_margin = Inches(.78)
section.right_margin = Inches(.78)
section.header_distance = Inches(.2)
section.footer_distance = Inches(.25)

normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string('000000')
normal.paragraph_format.line_spacing = Pt(16)
normal.paragraph_format.space_after = Pt(6)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)

for name, size, line, before, after in [('Title', 19.5, 25, 0, 6), ('Heading 1', 12.5, 18, 8, 4)]:
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.underline = False
    style.font.color.rgb = RGBColor.from_string('000000')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    style.paragraph_format.line_spacing = Pt(line)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)

doc.core_properties.title = '123Proxy 不限流量海外爬虫代理'
doc.core_properties.subject = '面向第三方爬虫平台的产品广告图文'
doc.core_properties.author = '123Proxy'
doc.core_properties.keywords = '海外爬虫代理,不限流量,隧道并发,YouTube,动态住宅,高带宽'
doc.core_properties.comments = ''


def paragraph(text='', style=None, size=None, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    run_font(p.add_run(text), size=size)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if size is not None and size < 11:
        p.paragraph_format.line_spacing = Pt(14)
    return p


def heading(text):
    return paragraph(text, style='Heading 1', after=4, keep=True)


def add_picture(p, source, width, alt):
    buffer = BytesIO()
    with Image.open(source) as image:
        image.convert('RGB').save(buffer, format='PNG')
    buffer.seek(0)
    shape = p.add_run().add_picture(buffer, width=Inches(width))
    shape._inline.docPr.set('descr', alt)
    return shape


def figure(filename, alt, width=6.5, after=4, keep=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Image paragraphs need auto spacing, not the exact body line height.
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    add_picture(p, ASSETS / filename, width, alt)
    return p


def hyperlink(p, text, url, size=10):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn(f'w:{key}'), FONT)
    properties.append(fonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), BLUE)
    properties.append(color)
    text_size = OxmlElement('w:sz')
    text_size.set(qn('w:val'), str(int(size * 2)))
    properties.append(text_size)
    run.append(properties)
    content = OxmlElement('w:t')
    content.text = text
    run.append(content)
    element.append(run)
    p._p.append(element)


def product_link(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    hyperlink(p, f'{label}  {url}', url)
    return p


# Page one introduces the article and the concurrency product. Only one hero.
paragraph('123Proxy 不限流量海外爬虫代理', style='Title', after=6, keep=True)
paragraph('面向爬虫开发工程师，提供隧道并发、不限量动态住宅和高带宽定制。网页爬虫、YouTube 下载与 AI 大规模下载，可按任务选择不限流量方案。', after=6, keep=True)
figure('01-main-banner.png', '123Proxy 不限流量海外爬虫代理 主 Banner', width=6.5, after=6, keep=False)

heading('持续网页爬虫与 YouTube 视频下载')
paragraph('常规 YouTube / yt-dlp 下载和持续网页爬虫，可先测试隧道代理按并发套餐。按同时在途请求配置容量，不按累计 GB 计费；兼容 requests、Scrapy 等代码接入。', after=6, keep=True)
figure('02-tunnel-plans.png', '隧道代理按并发套餐 折后月价 5线程270元 50线程2700元 200线程9000元')
paragraph('不限流量不等于不限速度，实际吞吐受线程容量、目标响应和网络状况影响。', size=9.5, after=3)
product_link('隧道套餐价格', 'https://www.123proxy.cn/pricing.html?product=tunnel')

doc.add_page_break()

# Page two explains two genuinely different products, rather than repeating a slogan.
heading('住宅浏览器自动化与持续任务')
paragraph('住宅浏览器任务流量较大时，可评估不限量动态住宅。按端口和周期购买，每端口不限流量与并发；实际吞吐仍受套餐带宽限制。', after=6, keep=True)
figure('03-residential-ports.png', '不限量动态住宅 5端口600元每月 每端口不限流量与并发 固定3至30分钟轮换 地区按套餐统一设置')
product_link('端口套餐价格', 'https://www.123proxy.cn/pricing.html?product=unlimited')

heading('AI 大规模下载与高带宽定制')
paragraph('视频、图片、公开代码和文档的海量下载，可配置高带宽代理 IP 与定制代理池。技术团队按目标站点、数据量与完成周期测试评估后报价，并独立交付。', after=6, keep=True)
figure('04-high-bandwidth-project.png', '高带宽代理IP 10Gbps以上单项目聚合方案需测试评估 公开数据源经定制代理池至爬虫集群和客户存储 按项目报价 技术支持独立交付')
product_link('高带宽项目咨询', 'https://www.123proxy.cn/contact.html#solutions')

doc.add_page_break()

# Keep the original ad's flow-priced alternative, while separating it from unlimited plans.
heading('国家定位与固定出口方案')
paragraph('需要逐请求指定国家或地区、保持住宅 SESSION，可选按流量购买的隧道住宅代理。需要长期固定出口，则选择长效静态代理或长效静态住宅。', after=6)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.45, 1.86, 3.63]
for cell, width, text in zip(table.rows[0].cells, widths, ['产品', '适用需求', '套餐价格示例']):
    cell.width = Inches(width)
    cell.text = text

rows = [
    ('隧道住宅代理', '国家或地区定位\n住宅 SESSION', f'1GB {money(price("residentialDynamicIp", 1))}\n10,000GB 档折合 ¥3/GB\n仅按流量购买，不属于不限流量套餐'),
    ('长效静态代理', '长期固定机房出口\n来源 IP 白名单', f'1 IP {money(price("fixedIp", 1))}/月\n1,000 IP 档折合 ¥20/IP/月'),
    ('长效静态住宅', '长期固定住宅出口\n稳定住宅身份', f'1 IP {money(price("residentialStaticIp", 1))}/月\n500 IP 档折合 ¥40/IP/月'),
]
assert price('residentialDynamicIp', 10000) / 10000 == 3
assert price('fixedIp', 1000) / 1000 == 20
assert price('residentialStaticIp', 500) / 500 == 40

for data in rows:
    cells = table.add_row().cells
    for cell, width, text in zip(cells, widths, data):
        cell.width = Inches(width)
        cell.text = text

borders = OxmlElement('w:tblBorders')
for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    border = OxmlElement(f'w:{side}')
    for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
        border.set(qn(f'w:{key}'), value)
    borders.append(border)
table._tbl.tblPr.append(borders)

for index, row in enumerate(table.rows):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    if index == 0:
        row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        properties = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for side, value in [('top', '100'), ('bottom', '100'), ('left', '110'), ('right', '110')]:
            margin = OxmlElement(f'w:{side}')
            margin.set(qn('w:w'), value)
            margin.set(qn('w:type'), 'dxa')
            margins.append(margin)
        properties.append(margins)
        if index == 0:
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), 'E7EFF9')
            properties.append(fill)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14.5)
            for run in p.runs:
                run_font(run, size=10, bold=index == 0)

paragraph('以上为对应规格的当前报价。批量折合价需购买对应档位，地区库存和最终金额以控制台确认为准。', size=9.5, after=6)

heading('先测试目标站点 再购买生产套餐')
paragraph('隧道并发套餐可申请 5 线程 4 小时测试；不限量动态住宅可申请 5 端口 2 小时测试。资格与开通方式由客服确认。高带宽项目由技术团队安排测试与方案评估。', after=5)
paragraph('咨询时请说明目标站点、爬虫框架、并发或下载量，以及住宅和地区需求。', after=5)

codes = doc.add_paragraph()
codes.alignment = WD_ALIGN_PARAGRAPH.CENTER
codes.paragraph_format.line_spacing = 1.0
codes.paragraph_format.space_before = Pt(3)
codes.paragraph_format.space_after = Pt(4)
codes.paragraph_format.keep_with_next = True
add_picture(codes, ASSETS / 'contact-service-wecom.png', 1.4, '官网客户服务企业微信二维码 免费测试和套餐办理')
run_font(codes.add_run('        '))
add_picture(codes, ASSETS / 'contact-solutions-wechat.png', 1.46, '官网技术与企业方案微信二维码 高带宽项目咨询')
caption = paragraph('免费测试与客户服务                    技术咨询与高带宽项目', size=9.5, after=5, keep=True)
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

product_link('官网与实时价格', 'https://www.123proxy.cn/pricing.html')
product_link('开发者接入手册', 'https://www.123proxy.cn/developers/')
paragraph('商务邮箱  sales@123proxy.cn', size=10, after=4)
paragraph('使用说明：爬虫程序需部署在海外网络环境，代理仅用于合法公开数据爬取等业务，不提供 VPN 或翻墙功能。', size=9.5, after=4)
paragraph('报价与产品规则核对日期为 2026 年 10 月 10 日，后续更新请查看官网。', size=9.5, after=0)

# Disable CJK document-grid snapping so physical image and paragraph spacing stays reliable.
all_paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
for p in all_paragraphs:
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)

doc.save(OUT)
print(json.dumps({'output': str(OUT), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'images': len(doc.inline_shapes)}, ensure_ascii=False))
